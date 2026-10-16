"""產生 Ken 財商教室 Excel 試算表（含即時公式）到 public/downloads/ken/。

用法：python3 scripts/build-ken-sheets.py
黃色格子為輸入，綠色格子為公式結果。通膨一律以 10% 起跳（資料驗證＋MAX(0.1, x)）。
"""
import sys
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'public' / 'downloads' / 'ken'
REVIEWED = '2026-10-16'
FONT = '微軟正黑體'
IN_FILL = PatternFill('solid', fgColor='FFF4CC')
OUT_FILL = PatternFill('solid', fgColor='E2EFDA')
HEAD_FILL = PatternFill('solid', fgColor='17160F')
SEC_FILL = PatternFill('solid', fgColor='E8DFCD')
THIN = Side(style='thin', color='C3B89F')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT = {'money': '#,##0', 'pct': '0.0%', 'num': '#,##0.0', 'int': '#,##0', 'date': 'yyyy-mm-dd', 'text': '@', 'gen': 'General'}
DISCLAIMER = f'本試算表為教育用途的估算工具，不構成投資、稅務、法律、保險或收入的建議與保證；法規數字以主管機關最新公告為準。資料檢視日：{REVIEWED}。'
INFL_NOTE = '個人／家庭通膨（10% 起跳）'


def f(size=12, bold=False, color='17160F'):
    return Font(name=FONT, size=size, bold=bold, color=color)


class Sheet:
    def __init__(self, wb, name, title, desc, widths=(34, 18, 16, 16, 16, 16, 16)):
        self.ws = wb.create_sheet(name)
        self.a = {}
        self.ws['A1'] = title
        self.ws['A1'].font = f(16, True)
        self.ws['A2'] = desc + '　黃色格子可修改，綠色格子會自動計算。'
        self.ws['A2'].font = f(12, color='5B5749')
        self.ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
        self.ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(widths))
        self.ws.row_dimensions[2].height = 48
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.r = 4

    def cell(self, row, col, value, kind='gen', fill=None, bold=False, size=12, wrap=False):
        c = self.ws.cell(row=row, column=col, value=value)
        c.font = f(size, bold)
        c.number_format = FMT[kind]
        if fill is not None:
            c.fill = fill
            c.border = BOX
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical='top')
        return c

    def ref(self, s):
        return s.format(**self.a)

    def section(self, label):
        self.r += 1
        c = self.cell(self.r, 1, label, bold=True, size=13)
        for col in range(1, len(self.ws.column_dimensions) + 1):
            self.ws.cell(row=self.r, column=col).fill = SEC_FILL
        c.fill = SEC_FILL
        self.r += 1

    def inp(self, key, label, value, kind='money', note=''):
        self.cell(self.r, 1, label, wrap=True)
        self.cell(self.r, 2, value, kind, IN_FILL)
        if note:
            self.cell(self.r, 3, note, size=11).font = f(11, color='5B5749')
        self.a[key] = f'$B${self.r}'
        self.r += 1
        return self.a[key]

    def inflation(self, key='i'):
        addr = self.inp(key, INFL_NOTE, 0.10, 'pct', '不可低於 10%')
        dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0.1', showErrorMessage=True,
                            errorTitle='通膨假設', error='本工具假設台灣個人通膨每年 10% 起跳，請輸入 10% 以上。')
        self.ws.add_data_validation(dv)
        dv.add(addr.replace('$', ''))
        self.a[key + 'x'] = f'MAX(0.1,{addr})'
        return addr

    def out(self, key, label, formula, kind='money', note='', bold=True):
        self.cell(self.r, 1, label, bold=bold, wrap=True)
        self.cell(self.r, 2, self.ref(formula), kind, OUT_FILL, bold=bold)
        if note:
            self.cell(self.r, 3, note, size=11).font = f(11, color='5B5749')
        self.a[key] = f'$B${self.r}'
        self.r += 1
        return self.a[key]

    def note(self, text):
        self.cell(self.r, 1, text, size=11, wrap=True).font = f(11, color='5B5749')
        self.ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=len(self.ws.column_dimensions))
        self.ws.row_dimensions[self.r].height = 34
        self.r += 1

    def header(self, labels):
        for i, h in enumerate(labels, 1):
            c = self.cell(self.r, i, h, bold=True, wrap=True)
            c.fill = HEAD_FILL
            c.font = f(12, True, 'FFFFFF')
            c.border = BOX
        self.ws.row_dimensions[self.r].height = 36
        self.r += 1
        return self.r

    def row(self, values, kinds):
        """values: 常數為輸入；以 '=' 開頭的字串為公式，可用 {r} 代表本列、其他 {key}。"""
        for i, (v, k) in enumerate(zip(values, kinds), 1):
            if isinstance(v, str) and v.startswith('='):
                self.cell(self.r, i, self.ref(v.replace('{r}', str(self.r))), k, OUT_FILL)
            else:
                self.cell(self.r, i, v, k, IN_FILL)
        self.r += 1
        return self.r - 1

    def validate_list(self, rng, options):
        dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
        self.ws.add_data_validation(dv)
        dv.add(rng)


def intro(wb, title, purpose, sheets):
    ws = wb.active
    ws.title = '說明'
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 90
    ws['A1'] = title
    ws['A1'].font = f(18, True)
    ws['A2'] = purpose
    ws['A2'].font = f(12)
    ws['A2'].alignment = Alignment(wrap_text=True)
    ws.merge_cells('A2:B2')
    ws.row_dimensions[2].height = 40
    r = 4
    for k, v in [('使用方式', '黃色格子填入自己的數字，綠色格子會自動計算；範例數字可直接覆蓋。'),
                 ('通膨假設', '台灣個人／家庭通膨每年 10% 起跳。官方 CPI 約 2% 左右，但房租、外食、醫療與照護的實際漲幅常遠高於 CPI；本工具採保守的 10% 起跳假設，讓規劃留足安全邊際。通膨格子設有驗證，低於 10% 會被拒絕，公式也以 MAX(10%, 輸入值) 計算。'),
                 ('線上版', 'https://ak-wealth-global.vercel.app/resources/ken/'),
                 ('免責聲明', DISCLAIMER)]:
        ws.cell(row=r, column=1, value=k).font = f(12, True)
        c = ws.cell(row=r, column=2, value=v)
        c.font = f(12)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 64 if len(v) > 60 else 24
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='工作表').font = f(13, True)
    r += 1
    for name, desc in sheets:
        ws.cell(row=r, column=1, value=name).font = f(12, True)
        c = ws.cell(row=r, column=2, value=desc)
        c.font = f(12)
        c.alignment = Alignment(wrap_text=True)
        r += 1


def years_table(s, first, last, step, cols, kinds):
    """cols: 以 {y} 代表年數的公式樣板。"""
    for y in range(first, last + 1, step):
        s.row([y] + [c.replace('{y}', str(y)) if isinstance(c, str) else c for c in cols], kinds)


# ======================================================================
def basics():
    wb = Workbook()
    intro(wb, 'Ken 財商基礎試算表', '先看懂自己的位置：時間換金錢的真相、個人通膨、預算比例、開支縮減與五年行動計畫。', [
        ('時薪真相', '淨資產 ÷ 總工時＝每工作 1 小時替自己留下多少錢，並對照台灣財富 10 等級。'),
        ('財務階段', '從財務安全到財務傳承，每一階需要多少資本（含通膨）與缺口。'),
        ('個人通膨', '用常買品項的價格算出自己的通膨與 20 年購買力。'),
        ('薪資分配', '6 罐理財食譜、631、532 三種分配法與目前實際花費比較。'),
        ('開支縮減', '逐項縮減支出，並把省下的錢投入 5 年的累積。'),
        ('存款vs資產', '定存與抗通膨資產 30 年的名目與實質購買力。'),
        ('行動計畫', '債務清償、收入型資產與夢想清單。')])

    s = Sheet(wb, '時薪真相', '時薪真相與財富等級', '如果今天全部變賣，扣掉所有負債，你工作至今每小時替自己留下多少淨資產？', (40, 20, 22, 46))
    s.section('工時')
    s.inp('y', '已工作年數（扣除失業、育嬰等）', 20, 'num')
    s.inp('w', '平均每週工時', 45, 'num')
    s.out('h', '工作至今總工時', '={y}*52*{w}', 'int')
    s.section('資產（變賣價值）')
    a0 = s.r
    for n, v in [('房地產', 12000000), ('汽機車', 400000), ('存款、股票、基金、保單價值', 1500000), ('家具家電、3C、收藏', 100000), ('其他', 0)]:
        s.row([n, v], ['text', 'money'])
    s.a['A'] = f'SUM(B{a0}:B{s.r - 1})'
    s.section('負債（尚欠金額）')
    d0 = s.r
    for n, v in [('房貸', 7000000), ('車貸', 200000), ('信用卡、分期', 0), ('信貸、學貸、親友借款', 0), ('其他', 0)]:
        s.row([n, v], ['text', 'money'])
    s.a['D'] = f'SUM(B{d0}:B{s.r - 1})'
    s.section('結果')
    s.out('assets', '資產合計', '={A}')
    s.out('debts', '負債合計', '={D}')
    s.out('net', '淨資產', '={assets}-{debts}')
    s.out('per', '每工作 1 小時留下的淨資產', '=IF({h}>0,{net}/{h},0)', 'money', '對照 2026 年基本工資時薪 196 元')
    s.inflation()
    s.out('real', '若淨資產不增加，10 年後實質購買力', '={net}/(1+{ix})^10')
    s.section('台灣財富淨資產 10 等級（由低到高排列供查表）')
    t0 = s.header(['等級', '淨資產下限', '生活樣貌'])
    tiers = [('第十等 赤貧', -9.99e15, '破產、嚴重債務'), ('第九等 貧困', -5000000, '負債累累，依賴補助'), ('第八等 低收入', 0, '低薪、不穩定收入'),
             ('第七等 中低產', 1000000, '一般受薪，無房產或高額房貸'), ('第六等 普通', 5000000, '穩定受薪、小型自營'), ('第五等 小康', 20000000, '專業人士、穩定雙薪'),
             ('第四等 中高產', 100000000, '中小企業主、資深經理人'), ('第三等 富裕', 500000000, '成功企業家、多處資產'), ('第二等 富豪', 2000000000, '集團創辦人、大股東'),
             ('第一等 極富裕', 10000000000, '私人飛機、家族基金會')]
    for n, m, life in tiers:
        s.cell(s.r, 1, n)
        s.cell(s.r, 2, m, 'money')
        s.cell(s.r, 3, life)
        s.r += 1
    t1 = s.r - 1
    s.r += 1
    s.out('tier', '你的財富等級', f'=INDEX(A{t0}:A{t1},MATCH({{net}},B{t0}:B{t1},1))', 'text')
    s.out('next', '距離下一等級還差', f'=IFERROR(INDEX(B{t0}:B{t1},MATCH({{net}},B{t0}:B{t1},1)+1)-{{net}},0)')

    s = Sheet(wb, '財務階段', '財務階段與資本缺口', '假設從今天起不再有工作收入：現有資本撐得了多久？每個財務階段需要多少資本？', (34, 20, 20, 20, 20, 16))
    s.inp('m', '每月基本生活開銷（今日幣值）', 40000)
    s.inp('d', '理想生活每月所需', 80000)
    s.inp('cap', '現在可動用的資本', 2000000)
    s.inp('n', '預計還會活幾年', 30, 'num')
    s.inp('save', '目前每年能存下', 200000)
    s.inp('leg', '想留給家人或公益（今日幣值）', 5000000)
    s.inflation()
    s.out('lastm', '資本可撐約幾個月（逐年通膨）', '=IFERROR(12*LN(1+{cap}*{ix}/({m}*12))/LN(1+{ix}),0)', 'num')
    s.section('五個財務階段（所需資本＝每年支出逐年隨通膨成長的總和）')
    s.header(['階段', '定義', '所需資本', '缺口', '照存錢速度需幾年', '是否達成'])
    g = '12*{M}*((1+{ix})^{N}-1)/{ix}'
    rows = [('財務安全', '36 個月基本開銷', g.replace('{M}', '{m}').replace('{N}', '3')),
            ('財務穩定', '基本開銷撐到終老', g.replace('{M}', '{m}').replace('{N}', '{n}')),
            ('財務獨立', '基本開銷 +20% 撐到終老', g.replace('{M}', '{m}*1.2').replace('{N}', '{n}')),
            ('財務自由', '理想生活撐到終老', g.replace('{M}', '{d}').replace('{N}', '{n}')),
            ('財務傳承', '理想生活＋傳承', g.replace('{M}', '{d}').replace('{N}', '{n}') + '+{leg}*(1+{ix})^{n}')]
    for name, desc, formula in rows:
        s.row([name, desc, '=' + formula, '=MAX(0,C{r}-{cap})', '=IF(D{r}=0,0,IF({save}>0,D{r}/{save},"無法達成"))', '=IF(D{r}=0,"✔ 已達成","未達成")'],
              ['text', 'text', 'money', 'money', 'num', 'text'])
    s.note('「照存錢速度」只算存錢、不含投資報酬，用來看清只靠工作儲蓄要多久。')

    s = Sheet(wb, '個人通膨', '個人通膨與購買力貶值', '官方 CPI 不等於你的生活成本。填入常買品項當年與現在的價格，算出自己的通膨。', (16, 22, 16, 16, 16, 16))
    s.inp('span', '價格比較間隔（年）', 5, 'num')
    s.inp('sal', '目前月薪', 50000)
    s.inflation()
    start = s.header(['類別', '品項', '當年價格', '現在價格', '漲幅', '年化通膨'])
    items = [('食物雜貨', '便當', 80, 120), ('食物雜貨', '雞蛋一盒', 60, 85), ('交通', '機車加滿油', 150, 190), ('生活必需', '套房月租', 12000, 16000),
             ('生活必需', '電費（月）', 900, 1200), ('生活必需', '健保門診部分負擔', 150, 200), ('娛樂社交', '外帶咖啡', 50, 70), ('娛樂社交', '電影票', 280, 350)]
    kinds = ['text', 'text', 'money', 'money', 'pct', 'pct']
    for c_, n, o, nw in items:
        s.row([c_, n, o, nw, '=IF(C{r}>0,D{r}/C{r}-1,"")', '=IF(C{r}>0,(D{r}/C{r})^(1/{span})-1,"")'], kinds)
    for _ in range(6):
        s.row([None, None, None, None, '=IF(C{r}>0,D{r}/C{r}-1,"")', '=IF(C{r}>0,(D{r}/C{r})^(1/{span})-1,"")'], kinds)
    end = s.r - 1
    s.validate_list(f'A{start}:A{end}', ['食物雜貨', '交通', '生活必需', '娛樂社交'])
    s.r += 1
    s.out('avg', '你的個人通膨（平均年化）', f'=IFERROR(AVERAGE(F{start}:F{end}),0)', 'pct')
    s.out('used', '規劃採用通膨（實測與 10% 取高）', '=MAX({ix},{avg})', 'pct')
    s.section('月薪購買力時間表')
    s.header(['第幾年', '實質購買力', '維持生活所需月薪'])
    for y in [1, 2, 3, 5, 10, 15, 20]:
        s.row([y, f'={{sal}}/(1+{{used}})^{y}', f'={{sal}}*(1+{{used}})^{y}'], ['int', 'money', 'money'])

    s = Sheet(wb, '薪資分配', '薪資自動分配（6 罐／631／532）', '發薪日輸入實拿薪資，自動拆出每一罐的金額；比例可自行修改，合計需為 100%。', (22, 14, 18, 18, 18, 14))
    s.inp('inc', '本月實拿薪資', 60000)
    s.inflation()
    for title, parts in [('6 罐理財食譜', [('生活開銷', .4), ('收入型資產', .2), ('自我教育', .1), ('玩樂夢想', .1), ('感恩回饋', .1), ('債務清償', .1)]),
                         ('631 法則', [('生活開銷', .6), ('投資理財', .3), ('緊急預備金', .1)]),
                         ('532 法則', [('必要生活', .5), ('想要／享受', .3), ('儲蓄投資', .2)])]:
        s.section(title)
        st = s.header(['用途', '比例', '本月金額', '目前實際花費', '差額（實際−目標）'])
        for n, p in parts:
            s.row([n, p, '={inc}*B{r}', 0, '=D{r}-C{r}'], ['text', 'pct', 'money', 'money', 'money'])
        s.row(['合計', f'=SUM(B{st}:B{s.r - 1})', f'=SUM(C{st}:C{s.r - 1})', f'=SUM(D{st}:D{s.r - 1})', f'=SUM(E{st}:E{s.r - 1})'], ['text', 'pct', 'money', 'money', 'money'])
        s.ws.cell(row=s.r - 1, column=1).value = '合計（比例需＝100%）'
    s.r += 1
    s.out('next', '明年同樣生活開銷（6 罐生活罐 × 通膨）', '={inc}*0.4*(1+{ix})')

    s = Sheet(wb, '開支縮減', '開支縮減計畫', '列出可以降低的支出與開始日期；省下的錢投入收入型資產，看 5 年後累積多少。', (30, 16, 16, 16, 16))
    s.inp('rate', '投資年報酬（假設）', 0.06, 'pct')
    s.inp('yrs', '持續年數', 5, 'num')
    s.inflation()
    st = s.header(['支出項目', '目前每月', '調整後每月', '每月省下', '開始日期'])
    for n, a, b in [('手機資費', 1399, 599), ('串流與訂閱', 990, 390), ('外食與飲料', 9000, 6500), ('保險（重複保障）', 5000, 3500), ('交通（改通勤方式）', 4000, 2800)]:
        s.row([n, a, b, '=B{r}-C{r}', None], ['text', 'money', 'money', 'money', 'date'])
    for _ in range(10):
        s.row([None, None, None, '=B{r}-C{r}', None], ['text', 'money', 'money', 'money', 'date'])
    en = s.r - 1
    s.r += 1
    s.out('m', '每月省下', f'=SUM(D{st}:D{en})')
    s.out('yr', '每年省下', '={m}*12')
    s.out('fv', '每月投入、持續後累積', '=FV({rate}/12,{yrs}*12,-{m})')
    s.out('real', '換算今日購買力', '={fv}/(1+{ix})^{yrs}')

    s = Sheet(wb, '存款vs資產', '存款 vs 抗通膨資產', '同一筆錢放定存或放進會增值的資產，名目金額與實質購買力差多少？資產增值率是假設值，過去表現不代表未來。', (12, 20, 20, 20, 20))
    s.inp('p', '本金', 1000000)
    s.inp('dep', '定存年利率', 0.017, 'pct')
    s.inp('asset', '資產年增值（假設，例如黃金長期平均）', 0.075, 'pct')
    s.inflation()
    s.header(['年', '定存名目', '定存實質', '資產名目', '資產實質'])
    years_table(s, 1, 30, 1, ['={p}*(1+{dep})^{y}', '={p}*((1+{dep})/(1+{ix}))^{y}', '={p}*(1+{asset})^{y}', '={p}*((1+{asset})/(1+{ix}))^{y}'], ['int'] + ['money'] * 4)

    s = Sheet(wb, '行動計畫', '債務清償與五年行動計畫', '先還高利率債務，再建立收入型資產；寫下每一個目標的金額與日期。', (28, 16, 12, 16, 16, 16))
    s.inflation()
    s.section('債務清償')
    st = s.header(['債務', '餘額', '年利率', '每月還款', '還清月數', '預計還清日'])
    for n, b, r, p in [('信用卡循環', 80000, .15, 5000), ('信貸', 300000, .06, 9000), ('車貸', 200000, .03, 6000)]:
        s.row([n, b, r, p, '=IFERROR(NPER(C{r}/12,-D{r},B{r}),"還款不足")', '=IFERROR(EDATE(TODAY(),ROUNDUP(E{r},0)),"—")'], ['text', 'money', 'pct', 'money', 'num', 'date'])
    for _ in range(4):
        s.row([None, None, None, None, '=IF(B{r}>0,IFERROR(NPER(C{r}/12,-D{r},B{r}),"還款不足"),"")', '=IF(B{r}>0,IFERROR(EDATE(TODAY(),ROUNDUP(E{r},0)),"—"),"")'], ['text', 'money', 'pct', 'money', 'num', 'date'])
    s.section('收入型資產目標')
    s.inp('pas', '目標每月被動收入（今日幣值）', 50000)
    s.inp('yl', '收入型資產年收益率（假設）', 0.05, 'pct')
    s.inp('yr', '希望幾年內達成', 5, 'num')
    s.inp('st', '目前已有可投資資本', 300000)
    s.out('tgt', '屆時目標月收入（含通膨）', '={pas}*(1+{ix})^{yr}')
    s.out('need', '需要的收入型資本', '={tgt}*12/{yl}')
    s.out('mon', '每月需投入', '=MAX(0,-PMT({yl}/12,{yr}*12,-{st},{need}))')
    s.section('夢想清單（玩樂）與感恩計畫')
    s.header(['項目', '總金額（今日幣值）', '目標月數', '含通膨總金額', '每月需存'])
    for n, a, m in [('日本家族旅行', 200000, 12), ('捐助弱勢兒童', 36000, 12)]:
        s.row([n, a, m, '=B{r}*(1+{ix})^(C{r}/12)', '=D{r}/C{r}'], ['text', 'money', 'int', 'money', 'money'])
    for _ in range(4):
        s.row([None, None, None, '=IF(C{r}>0,B{r}*(1+{ix})^(C{r}/12),"")', '=IF(C{r}>0,D{r}/C{r},"")'], ['text', 'money', 'int', 'money', 'money'])
    s.note('12–24 個月、24–48 個月、48–60 個月各寫下一個具體目標，每月檢視一次。')
    return wb


# ======================================================================
def seniors():
    wb = Workbook()
    intro(wb, 'Ken 熟齡退休試算表', '中老年人（熟齡族、退休族）的核心需求：資產保全、穩定現金流、醫療風險控管。字體已放大，方便列印。', [
        ('退休提領', '整合勞保年金、勞退、股息與存款，算出安全提領額與資產耗盡年齡。'),
        ('長照醫療', '看護費、耗材、自費醫療在通膨下的總額與保險缺口。'),
        ('保單總表', '全家保單明細、年繳保費與期滿日。'),
        ('簡易收支', '只分三類：日常基本開銷、醫療保健、休閒娛樂。'),
        ('遺產贈與稅', '115 年標準的遺產稅與贈與稅，比較一次繼承與分年贈與。')])

    s = Sheet(wb, '退休提領', '退休金安全提領率', '每年需提領＝生活費（逐年通膨）－固定收入；資產每年先提領再以報酬率成長。', (16, 18, 18, 18, 18, 18))
    s.inp('age', '現在年齡', 65, 'int')
    s.inp('to', '規劃到幾歲（預期壽命再加 5–10 年）', 90, 'int')
    s.inp('m', '每月生活費（今日幣值）', 50000)
    s.inp('A', '存款與投資總額', 5000000)
    s.inp('ret', '資產年報酬（保守估）', 0.04, 'pct')
    s.inflation()
    s.section('每月固定收入')
    s.inp('ins', '勞保平均月投保薪資', 45800, 'money', '目前最高 45,800')
    s.inp('sen', '勞保年資', 30, 'num')
    s.out('est', '勞保老年年金估算（月）', '=MAX({ins}*{sen}*0.775%+3000,{ins}*{sen}*1.55%)', 'money', '實際以勞保局試算為準 calcr2.mol.gov.tw/trial')
    s.inp('lab', '勞保年金月領（填 0 用上方估算）', 0)
    s.inp('pen', '勞退個人專戶月領', 8000)
    s.inp('div', '每年股息、租金等固定收入', 80000, 'money', '產生股息的持股不要重複算進資產')
    s.out('fix', '每年固定收入合計', '=(IF({lab}>0,{lab},{est})+{pen})*12+{div}')
    s.section('結果')
    n_rows = 46
    t0 = s.r + 8  # 表格第一列資料列
    s.out('safe', '首年安全提領額（逐年隨通膨調升）', f'={{A}}/SUMPRODUCT(((1+{{ix}})/(1+{{ret}}))^(ROW(INDIRECT("1:"&MAX(1,{{to}}-{{age}})))-1))')
    s.out('rate', '安全提領率', '={safe}/{A}', 'pct')
    s.out('rule4', '4% 法則首年可提領（僅供對照）', '={A}*4%', 'money', '4% 法則以 2–3% 通膨為前提')
    s.out('mon', '每月安全總支出（固定收入＋提領）', '=({fix}+{safe})/12')
    s.out('dead', '資產耗盡年齡', f'=IFERROR(INDEX(A{t0}:A{t0 + n_rows - 1},MATCH(TRUE,INDEX(F{t0}:F{t0 + n_rows - 1}<0,0),0)),"規劃期間內未耗盡")', 'gen')
    s.out('gap', '資金缺口（今日需補足）', f'=MAX(0,SUMPRODUCT(D{t0}:D{t0 + n_rows - 1}/(1+{{ret}})^(A{t0}:A{t0 + n_rows - 1}-{{age}}))-{{A}})')
    s.r += 1
    s.header(['年齡', '年支出', '固定收入', '需提領', '期初資產', '期末資產'])
    assert s.r == t0, (s.r, t0)
    for k in range(n_rows):
        if k == 0:
            s.row(['={age}', '={m}*12', '={fix}', '=IF(A{r}<{to},MAX(0,B{r}-C{r}),0)', '={A}', '=IF(A{r}<{to},(E{r}-D{r})*(1+{ret}),"")'], ['int', 'money', 'money', 'money', 'money', 'money'])
        else:
            p = s.r - 1
            s.row([f'=A{p}+1', f'=B{p}*(1+{{ix}})', '={fix}', '=IF(A{r}<{to},MAX(0,B{r}-C{r}),0)', f'=IF(A{{r}}<{{to}},F{p},"")', '=IF(A{r}<{to},(E{r}-D{r})*(1+{ret}),"")'], ['int', 'money', 'money', 'money', 'money', 'money'])
    s.note('固定收入以不調整計算；勞保年金實際會依 CPI 累計成長 5% 調整。')

    s = Sheet(wb, '長照醫療', '長照與重大醫療現金流評估', '未來照護花費以通膨推估；保險給付多為固定金額，不隨通膨增加。', (40, 20, 40))
    s.inp('u', '距離可能需要照護（年）', 15, 'num')
    s.inp('y', '照護年數', 7.3, 'num', '衛福部常引用國人平均長照需求約 7.3 年')
    s.inp('care', '看護費每月（今日幣值）', 50000, 'money', '外籍看護約 3 萬起，居家或機構常見 4–6 萬')
    s.inp('sup', '耗材、營養品每月', 8000)
    s.inp('med', '一次性自費醫療（標靶、自費醫材）', 1000000)
    s.inp('home', '居家無障礙改裝、輔具', 200000)
    s.inflation()
    s.section('保險與補助')
    s.inp('lump', '長照／失能險一次給付', 500000)
    s.inp('mp', '長照／失能險每月給付', 20000)
    s.inp('crit', '重大傷病、癌症險一次給付', 300000)
    s.inp('gov', '長照 2.0 等補助每月（估）', 10000, 'money', '依失能等級與資格，可撥 1966 查詢')
    s.inp('rr', '準備金投資年報酬', 0.04, 'pct')
    s.section('結果')
    s.out('start', '開始照護時每月總花費', '=({care}+{sup})*(1+{ix})^{u}')
    s.out('carec', '照護期間看護＋耗材合計', '=({care}+{sup})*(1+{ix})^{u}*((1+{ix})^(ROUND({y}*12,0)/12)-1)/((1+{ix})^(1/12)-1)')
    s.out('once', '自費醫療＋居家改裝（未來值）', '=({med}+{home})*(1+{ix})^{u}')
    s.out('need', '未來總需求', '={carec}+{once}')
    s.out('today', '換算今日幣值', '={need}/(1+{ix})^{u}')
    s.out('cover', '保險理賠＋補助合計', '={lump}+{crit}+({mp}+{gov})*ROUND({y}*12,0)')
    s.out('gap', '資金缺口', '=MAX(0,{need}-{cover})')
    s.out('save', '從今天起每月需準備', '=IF({u}>0,PMT({rr}/12,{u}*12,0,-{gap}),{gap})')

    s = Sheet(wb, '保單總表', '全家保單總表', '集中記錄全家保單，突發疾病時一眼看清能動用的資源。', (12, 28, 18, 14, 16, 14, 14, 18))
    s.inp('inc', '家庭年收入', 1200000)
    st = s.header(['被保險人', '保險公司／保單名稱', '類型', '年繳保費', '保額或日額', '繳費期滿', '保障期滿', '理賠電話'])
    k = ['text', 'text', 'text', 'money', 'money', 'date', 'date', 'text']
    for row in [('爸爸', '○○人壽 終身醫療', '醫療', 28000, 2000, date(2030, 6, 30), date(2060, 6, 30), ''), ('爸爸', '○○人壽 長照險', '長照／失能', 36000, 20000, date(2032, 3, 31), date(2070, 3, 31), ''),
                ('媽媽', '○○產險 意外險', '意外', 3500, 3000000, date(2027, 1, 31), date(2027, 1, 31), ''), ('媽媽', '○○人壽 癌症險', '癌症／重大傷病', 15000, 300000, date(2029, 12, 31), date(2050, 12, 31), '')]:
        s.row(list(row), k)
    for _ in range(16):
        s.row([None] * 8, k)
    en = s.r - 1
    types = ['壽險', '醫療', '癌症／重大傷病', '意外', '長照／失能', '儲蓄／年金']
    s.validate_list(f'C{st}:C{en}', types)
    s.r += 1
    s.out('tot', '每年保費合計', f'=SUM(D{st}:D{en})')
    s.out('ratio', '保費占家庭年收入', '=IF({inc}>0,{tot}/{inc},0)', 'pct', '常見參考：10–15% 以內')
    s.out('soon', '一年內繳費期滿張數', f'=COUNTIFS(F{st}:F{en},">="&TODAY(),F{st}:F{en},"<"&TODAY()+365)', 'int')
    s.section('依類型統計')
    s.header(['類型', '張數', '年繳保費'])
    for t in types:
        s.row([t, f'=COUNTIF(C{st}:C{en},A{{r}})', f'=SUMIF(C{st}:C{en},A{{r}},D{st}:D{en})'], ['text', 'int', 'money'])

    s = Sheet(wb, '簡易收支', '簡易生活收支表（長輩版）', '每花一筆錢記一列，類別用下拉選單選三類之一。', (16, 18, 32, 16, 16, 14))
    s.section('每月預算')
    b0 = s.r
    for n, v in [('日常基本開銷', 25000), ('醫療保健', 6000), ('休閒娛樂', 5000)]:
        s.row([n, v, f'=SUMIF($B$__S:$B$__E,A{{r}},$D$__S:$D$__E)', '=B{r}-C{r}'], ['text', 'money', 'money', 'money'])
    s.ws.cell(row=b0 - 1, column=2, value='預算').font = f(12, True)
    s.ws.cell(row=b0 - 1, column=3, value='已花').font = f(12, True)
    s.ws.cell(row=b0 - 1, column=4, value='剩餘').font = f(12, True)
    s.inflation()
    s.out('tot', '本月已花', f'=SUM(C{b0}:C{b0 + 2})')
    s.out('left', '本月還可以花', f'=SUM(B{b0}:B{b0 + 2})-{{tot}}')
    s.out('next', '明年同樣預算需要', f'=SUM(B{b0}:B{b0 + 2})*(1+{{ix}})')
    s.r += 1
    st = s.header(['日期', '類別', '說明', '金額'])
    k = ['date', 'text', 'text', 'money']
    for d_, c_, n, a in [(date(2026, 10, 1), '日常基本開銷', '菜市場', 1200), (date(2026, 10, 2), '醫療保健', '門診與藥', 450), (date(2026, 10, 5), '休閒娛樂', '和老同學聚餐', 800), (date(2026, 10, 6), '日常基本開銷', '水電瓦斯', 2100)]:
        s.row([d_, c_, n, a], k)
    for _ in range(80):
        s.row([None] * 4, k)
    en = s.r - 1
    s.validate_list(f'B{st}:B{en}', ['日常基本開銷', '醫療保健', '休閒娛樂'])
    for rr in range(b0, b0 + 3):
        c = s.ws.cell(row=rr, column=3)
        c.value = c.value.replace('__S', str(st)).replace('__E', str(en))
    s.ws.freeze_panes = f'A{st}'

    s = Sheet(wb, '遺產贈與稅', '遺產稅與贈與稅試算（115 年發生案件適用）', '遺產以土地公告現值、房屋評定標準價格計算，通常低於市價。', (40, 20, 44))
    s.section('遺產')
    for k_, l, v in [('land', '土地（公告土地現值）', 12000000), ('house', '房屋（評定標準價格）', 1500000), ('cash', '現金、存款', 8000000), ('stock', '股票、基金（時價）', 6000000), ('oth', '其他財產', 0), ('debt', '未償債務', 0)]:
        s.inp(k_, l, v)
    s.inp('sp', '配偶在世（1＝是，0＝否）', 1, 'int')
    s.inp('kids', '直系血親卑親屬人數', 2, 'int', '未成年者另依距成年年數加扣，本表未計')
    s.inp('par', '在世父母人數', 0, 'int')
    s.inp('dis', '身心障礙繼承人數', 0, 'int')
    s.out('gross', '遺產總額', '={land}+{house}+{cash}+{stock}+{oth}')
    s.out('ded', '免稅額＋扣除額', '=13330000+{sp}*5530000+{kids}*560000+{par}*1380000+1380000+{dis}*6930000+{debt}', 'money', '免稅 1,333 萬；喪葬費 138 萬')
    s.out('net', '課稅遺產淨額', '=MAX(0,{gross}-{ded})')
    s.out('tax', '預估遺產稅', '=IF({net}<=56210000,{net}*10%,IF({net}<=112420000,5621000+({net}-56210000)*15%,14052500+({net}-112420000)*20%))', 'money', '10%／15%／20%')
    s.section('生前贈與')
    s.inp('T', '打算贈與的總額', 10000000)
    s.inp('dn', '贈與人數（父母各自贈與＝2）', 2, 'int')
    s.inp('gy', '分幾年贈與', 3, 'int')
    gt = 'IF({X}<=0,0,IF({X}<=28110000,{X}*10%,IF({X}<=56210000,2811000+({X}-28110000)*15%,7026000+({X}-56210000)*20%)))'
    s.out('once', '一年內全部贈與的贈與稅', '={dn}*' + gt.replace('{X}', '({T}/{dn}-2440000)'), 'money', '每位贈與人每年免稅 244 萬')
    s.out('spread', '分年贈與的贈與稅合計', '={gy}*{dn}*' + gt.replace('{X}', '({T}/{dn}/{gy}-2440000)'))
    et = 'IF({X}<=56210000,{X}*10%,IF({X}<=112420000,5621000+({X}-56210000)*15%,14052500+({X}-112420000)*20%))'
    s.out('after', '贈與後剩餘遺產的遺產稅', '=' + et.replace('{X}', 'MAX(0,{gross}-{T}-{ded})'))
    s.out('total', '分年贈與＋其餘繼承總稅負', '={spread}+{after}')
    s.note('死亡前 2 年內贈與配偶、子女等的財產仍要併入遺產課稅。不動產移轉另有土地增值稅、契稅、房地合一稅，請洽地政士、會計師或國稅局。')
    return wb


# ======================================================================
def white():
    wb = Workbook()
    intro(wb, 'Ken 白領上班族試算表', '核心需求：薪資放大、債務清償、資產累積與節稅。', [
        ('FIRE路徑', '儲蓄率、報酬率與通膨下，何時達到財務獨立。'),
        ('XIRR報酬', '跨帳戶投入與贖回，算出真實年化報酬。'),
        ('貸款還本', '不同利率的月付金，以及每月多還本金能縮短多久、省多少利息。'),
        ('綜所稅', '115 年度所得（116 年 5 月申報）的簡化估算，含自提勞退節稅效果。'),
        ('再平衡', '目前市值與目標比例，算出應買應賣金額。')])

    s = Sheet(wb, 'FIRE路徑', 'FIRE 財務獨立路徑', '資產＝前一年資產×(1+報酬)＋當年收入－當年支出；目標本金＝當年支出 ÷ 提領率。', (12, 18, 18, 18, 18, 14))
    s.inp('age', '現在年齡', 32, 'int')
    s.inp('A', '目前投資資產', 1500000)
    s.inp('inc', '稅後年收入', 1000000)
    s.inp('g', '年收入成長率', 0.03, 'pct')
    s.inp('exp', '年生活支出', 600000)
    s.inp('ret', '投資年報酬（假設）', 0.07, 'pct')
    s.inp('wr', '退休後提領率', 0.04, 'pct')
    s.inflation()
    t0 = s.r + 5
    s.out('fi', '財務獨立年齡', f'=IFERROR(INDEX(A{t0}:A{t0 + 50},MATCH("達成",F{t0}:F{t0 + 50},0)),"50 年內無法達成")', 'gen')
    s.out('sr', '目前儲蓄率', '=1-{exp}/{inc}', 'pct')
    s.out('real', '實質報酬率', '=(1+{ret})/(1+{ix})-1', 'pct', '低於 0 代表資產實質縮水')
    s.r += 1
    s.header(['年齡', '年收入', '年支出', '目標本金', '年初資產', '狀態'])
    assert s.r == t0
    for k in range(51):
        p = s.r - 1
        if k == 0:
            s.row(['={age}', '={inc}', '={exp}', '=C{r}/{wr}', '={A}', '=IF(E{r}>=D{r},"達成","")'], ['int', 'money', 'money', 'money', 'money', 'text'])
        else:
            s.row([f'=A{p}+1', f'=B{p}*(1+{{g}})', f'=C{p}*(1+{{ix}})', '=C{r}/{wr}', f'=E{p}*(1+{{ret}})+B{p}-C{p}', '=IF(E{r}>=D{r},"達成","")'], ['int', 'money', 'money', 'money', 'money', 'text'])

    s = Sheet(wb, 'XIRR報酬', '資產配置真實年化報酬（XIRR）', '投入填負數、贖回或配息填正數；最後一列填今日總市值（正數）。', (16, 30, 18, 4, 40, 18))
    st = s.header(['日期', '帳戶／標的', '現金流'])
    k = ['date', 'text', 'money']
    for d_, n, a in [(date(2023, 1, 10), '台股 0050 定期定額', -120000), (date(2023, 7, 10), '美股 VT', -100000), (date(2024, 1, 10), '台股 0050 定期定額', -120000), (date(2024, 8, 20), '0056 配息（領回）', 8000),
                     (date(2025, 1, 10), '美元定存轉 ETF', -100000), (date(2025, 9, 1), '加密貨幣', -50000), (date(2026, 1, 10), '台股 0050 定期定額', -60000), (date(2026, 10, 16), '今日總市值', 620000)]:
        s.row([d_, n, a], k)
    en = s.r - 1
    s.note('新增紀錄請在「今日總市值」那一列上方插入列，公式範圍會自動延伸。')
    s.ws.cell(row=st, column=5, value='XIRR 年化報酬').font = f(12, True)
    c = s.cell(st, 6, f'=XIRR(C{st}:C{en},A{st}:A{en})', 'pct', OUT_FILL, True)
    s.ws.cell(row=st + 1, column=5, value=INFL_NOTE).font = f(12)
    s.cell(st + 1, 6, 0.10, 'pct', IN_FILL)
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0.1', showErrorMessage=True, error='本工具假設台灣個人通膨每年 10% 起跳，請輸入 10% 以上。')
    s.ws.add_data_validation(dv)
    dv.add(f'F{st + 1}')
    s.ws.cell(row=st + 2, column=5, value='扣除通膨後實質年報酬').font = f(12, True)
    s.cell(st + 2, 6, f'=(1+F{st})/(1+MAX(0.1,F{st + 1}))-1', 'pct', OUT_FILL, True)
    s.ws.cell(row=st + 3, column=5, value='累計投入').font = f(12)
    s.cell(st + 3, 6, f'=-SUMIF(C{st}:C{en},"<0")', 'money', OUT_FILL)
    s.ws.cell(row=st + 4, column=5, value='市值＋已領回').font = f(12)
    s.cell(st + 4, 6, f'=SUMIF(C{st}:C{en},">0")', 'money', OUT_FILL)
    s.ws.cell(row=st + 5, column=5, value='總損益').font = f(12, True)
    s.cell(st + 5, 6, f'=F{st + 4}-F{st + 3}', 'money', OUT_FILL, True)

    s = Sheet(wb, '貸款還本', '貸款攤還與提前還本', '本息平均攤還；每月額外多還本金可以縮短期限、節省利息。', (16, 18, 18, 18, 18, 18, 18))
    s.inp('P', '貸款金額', 10000000)
    s.inp('rate', '年利率', 0.022, 'pct')
    s.inp('yrs', '貸款年數', 30, 'int')
    s.inp('ex', '每月額外多還本金', 5000)
    s.out('pmt', '每月本息攤還', '=PMT({rate}/12,{yrs}*12,-{P})')
    s.out('int0', '原本總利息', '={pmt}*{yrs}*12-{P}')
    s.out('n1', '多還本金後還清月數', '=NPER({rate}/12,-({pmt}+{ex}),{P})', 'num')
    s.out('short', '縮短年數', '=({yrs}*12-{n1})/12', 'num')
    s.out('int1', '多還本金後總利息', '=({pmt}+{ex})*{n1}-{P}')
    s.out('saved', '省下利息', '={int0}-{int1}')
    s.section('不同利率月付金比較')
    s.header(['年利率', '月付金', '總利息'])
    for d in (-0.004, -0.002, 0, 0.002, 0.004, 0.008):
        s.row([f'={{rate}}+{d}', '=PMT(A{r}/12,{yrs}*12,-{P})', '=B{r}*{yrs}*12-{P}'], ['pct', 'money', 'money'])
    s.section('攤還表（左：原計畫；右：每月多還本金）')
    s.header(['期數', '原計畫利息', '原計畫本金', '原計畫餘額', '多還利息', '多還本金', '多還餘額'])
    for k in range(1, 361):
        p = s.r - 1
        prev_a = '{P}' if k == 1 else f'D{p}'
        prev_b = '{P}' if k == 1 else f'G{p}'
        s.row([k, f'=IF({prev_a}>0,{prev_a}*{{rate}}/12,0)', f'=IF({prev_a}>0,MIN({prev_a},{{pmt}}-B{{r}}),0)', f'=MAX(0,{prev_a}-C{{r}})',
               f'=IF({prev_b}>0,{prev_b}*{{rate}}/12,0)', f'=IF({prev_b}>0,MIN({prev_b},{{pmt}}+{{ex}}-E{{r}}),0)', f'=MAX(0,{prev_b}-F{{r}})'], ['int'] + ['money'] * 6)

    s = Sheet(wb, '綜所稅', '綜合所得稅節稅試算（115 年度所得，116 年 5 月申報）', '簡化估算：未含分開計稅、股利分開課稅、最低稅負制與房租扣除排富條件。', (44, 20, 46))
    s.inp('s1', '本人薪資收入', 900000)
    s.inp('s2', '配偶薪資收入', 0)
    s.inp('mar', '有配偶合併申報（1＝是）', 0, 'int')
    s.inp('oth', '其他所得（利息、股利、租金等）', 50000)
    s.inp('sav', '其中：利息等儲蓄投資所得', 20000)
    s.inp('dep', '扶養親屬（未滿 70 歲）', 0, 'int')
    s.inp('d70', '扶養 70 歲以上直系尊親屬', 0, 'int')
    s.inp('item', '列舉扣除合計（保險費、醫藥、捐贈、房貸利息）', 0)
    s.inp('rent', '房屋租金支出', 120000, 'money', '特別扣除上限 18 萬；採列舉時本表不計')
    s.inp('ltc', '長照特別扣除人數', 0, 'int')
    s.inp('kid', '6 歲以下子女', 0, 'int')
    s.inp('pct', '勞退自提比例（0–6%）', 0.06, 'pct')
    s.inp('basic', '每人基本生活費（待財政部公告）', 213000, 'money', '114 年度為 21.3 萬')
    s.section('計算')
    s.out('gross', '綜合所得總額（扣除自提勞退）', '=({s1}+{s2})*(1-MIN(6%,{pct}))+{oth}')
    s.out('ex', '免稅額', '=101000*(1+{mar}+{dep})+151500*{d70}', 'money', '每人 10.1 萬；70 歲以上 15.15 萬')
    s.out('gen', '一般扣除額（標準與列舉取高）', '=MAX(IF({mar}=1,272000,136000),{item})')
    s.out('sal', '薪資特別扣除', '=MIN({s1}*(1-MIN(6%,{pct})),227000)+MIN({s2}*(1-MIN(6%,{pct})),227000)')
    s.out('spec', '儲蓄投資＋長照＋幼兒學前＋房租', '=MIN({sav},270000)+{ltc}*180000+{kid}*150000+IF({item}>0,0,MIN({rent},180000))')
    s.out('gap', '基本生活費差額', '=MAX(0,(1+{mar}+{dep}+{d70})*{basic}-({ex}+{gen}+MIN({sav},270000)+{ltc}*180000+{kid}*150000))')
    s.out('net', '綜合所得淨額', '=MAX(0,{gross}-{ex}-{gen}-{sal}-{spec}-{gap})')
    bk = 'IF({X}<=610000,{X}*5%,IF({X}<=1380000,{X}*12%-42700,IF({X}<=2770000,{X}*20%-153100,IF({X}<=5190000,{X}*30%-430100,{X}*40%-949100))))'
    s.out('tax', '預估應納稅額', '=' + bk.replace('{X}', '{net}'))
    s.out('net0', '若不自提勞退的所得淨額', '=MAX(0,{net}+({s1}+{s2})*MIN(6%,{pct})-(MIN({s1},227000)+MIN({s2},227000)-{sal}))')
    s.out('tax0', '若不自提勞退的應納稅額', '=' + bk.replace('{X}', '{net0}'))
    s.out('save', '自提勞退省下的稅', '={tax0}-{tax}')
    s.note('級距：61 萬以下 5%｜138 萬以下 12%（累進差額 42,700）｜277 萬以下 20%（153,100）｜519 萬以下 30%（430,100）｜以上 40%（949,100）。正式數字以財政部公告為準。')

    s = Sheet(wb, '再平衡', '投資組合再平衡', '目標比例合計需為 100%；可先用新投入資金補足偏低的類別。', (22, 18, 14, 14, 18, 20))
    s.inp('add', '本次新投入資金', 30000)
    st = s.header(['資產', '目前市值', '目前比例', '目標比例', '目標金額', '應買（+）／賣（−）'])
    rows = [('台股 ETF', 420000, .4), ('美股／全球 ETF', 260000, .3), ('債券 ETF', 80000, .15), ('加密貨幣', 90000, .05), ('現金／定存', 60000, .1), (None, None, None), (None, None, None)]
    en = st + len(rows) - 1
    for n, v, t in rows:
        s.row([n, v, f'=IF(B{{r}}="","",B{{r}}/SUM(B{st}:B{en}))', t, f'=IF(D{{r}}="","",(SUM(B{st}:B{en})+{{add}})*D{{r}})', '=IF(D{r}="","",E{r}-B{r})'], ['text', 'money', 'pct', 'pct', 'money', 'money'])
    s.r += 1
    s.out('chk', '目標比例合計', f'=SUM(D{st}:D{en})', 'pct', '需＝100%')
    return wb


# ======================================================================
def blue():
    wb = Workbook()
    intro(wb, 'Ken 藍領職人試算表', '核心需求：加班權益、生存底線（緊急預備金）、技能轉型。', [
        ('加班費', '依勞基法計算平日、休息日、國定假日出勤的加班費，並檢查每月 46 小時上限。'),
        ('預備金水庫', '把薪水拆成生活必需、債務與儲蓄，算出 3–6 個月的生存底線。'),
        ('證照投資', '比較證照的學費、時間成本與加薪，算出回本月數。')])

    s = Sheet(wb, '加班費', '加班費與工時權益核算', '時薪＝月薪 ÷ 240。實際以勞動契約、工資清冊與勞動部試算 calcr2.mol.gov.tw 為準。', (16, 18, 12, 18, 40))
    s.inp('M', '月薪（時薪制填 0）', 36000)
    s.inp('H0', '時薪（月薪制可留 0）', 0)
    s.out('H', '換算時薪', '=IF({H0}>0,{H0},{M}/240)', 'num', '2026 年基本工資：月 29,500、時 196')
    s.out('min', '是否低於基本工資', '=IF({H0}>0,IF({H0}<196,"⚠ 時薪低於 196 元","符合"),IF({M}<29500,"⚠ 月薪低於 29,500 元","符合"))', 'text')
    s.r += 1
    st = s.header(['日期', '類型', '加班時數', '加班費'])
    form = ('=IF(C{r}="",0,IF(B{r}="平日加班",{H}*(MIN(C{r},2)*4/3+MAX(0,MIN(C{r},4)-2)*5/3),'
            'IF(B{r}="休息日出勤",{H}*(MIN(C{r},2)*4/3+MAX(0,MIN(C{r},8)-2)*5/3+MAX(0,MIN(C{r},12)-8)*8/3),'
            'IF(B{r}="國定假日出勤",{H}*(8+MAX(0,MIN(C{r},10)-8)*4/3+MAX(0,MIN(C{r},12)-10)*5/3),0))))')
    k = ['date', 'text', 'num', 'money']
    for d_, t, h in [(date(2026, 10, 2), '平日加班', 2), (date(2026, 10, 6), '平日加班', 3), (date(2026, 10, 10), '國定假日出勤', 8), (date(2026, 10, 17), '休息日出勤', 6), (date(2026, 10, 21), '平日加班', 4)]:
        s.row([d_, t, h, form], k)
    for _ in range(26):
        s.row([None, None, None, form], k)
    en = s.r - 1
    s.validate_list(f'B{st}:B{en}', ['平日加班', '休息日出勤', '國定假日出勤'])
    s.r += 1
    s.out('tot', '本月法定加班費', f'=SUM(D{st}:D{en})')
    s.out('hrs', '延長工時合計（平日＋休息日）', f'=SUMIF(B{st}:B{en},"平日加班",C{st}:C{en})+SUMIF(B{st}:B{en},"休息日出勤",C{st}:C{en})', 'num')
    s.out('cap', '是否超過每月 46 小時', '=IF({hrs}>46,"⚠ 超過上限","未超過")', 'text', '經勞資協議可放寬至 54 小時，每 3 個月 138 小時')
    s.note('平日延長：前 2 小時 ×4/3、第 3–4 小時 ×5/3｜休息日：前 2 小時 ×4/3、3–8 小時 ×5/3、9–12 小時 ×8/3｜國定假日：8 小時內加發一日工資，超過依延長工時加給。')

    s = Sheet(wb, '預備金水庫', '家庭硬性支出與緊急預備金水庫', '預備金放在與日常帳戶分開、可隨時提領的帳戶；存滿前不投資高風險商品。', (34, 18, 30))
    s.inp('inc', '每月實拿收入', 38000)
    s.section('每月必要支出')
    for k_, l, v in [('a', '房租／房貸', 10000), ('b', '伙食', 9000), ('c', '交通', 2500), ('d', '保險', 2000), ('e', '水電、電信、家用', 3500), ('f', '孝親、子女', 3000)]:
        s.inp(k_, l, v)
    s.inp('debt', '債務每月攤還', 3000)
    s.inp('saved', '目前已存預備金', 30000)
    s.inp('tm', '目標月數（建議 3–6）', 6, 'int')
    s.inflation()
    s.section('結果')
    s.out('need', '每月生存底線', '={a}+{b}+{c}+{d}+{e}+{f}+{debt}')
    s.out('cover', '目前預備金可撐（月）', '=IF({need}>0,{saved}/{need},0)', 'num')
    s.out('goal3', '3 個月預備金', '={need}*3')
    s.out('goal', '目標預備金', '={need}*{tm}')
    s.out('left', '每月可存', '={inc}-{need}')
    s.out('months', '存滿需要（月）', '=IF({left}>0,MAX(0,{goal}-{saved})/{left},"收支已透支，先縮減支出")', 'num')
    s.out('next', '明年生存底線（含通膨）', '={need}*(1+{ix})')
    s.out('p1', '薪水拆分：生活必需占比', '=({need}-{debt})/{inc}', 'pct')
    s.out('p2', '薪水拆分：債務占比', '={debt}/{inc}', 'pct')
    s.out('p3', '薪水拆分：儲蓄占比', '=MAX(0,{left})/{inc}', 'pct')

    s = Sheet(wb, '證照投資', '技能轉型與考照投資回報', '扣通膨淨效益：把未來的加薪折算成今天的購買力。可先查勞動力發展署訓練補助。', (28, 16, 12, 16, 16, 14, 18, 20))
    s.inp('Y', '評估年數', 10, 'int')
    s.inflation()
    s.header(['證照', '學費與報名費', '訓練時數', '請假少領收入', '每月加薪', '回本月數', '名目淨效益', '扣通膨淨效益'])
    k = ['text', 'money', 'num', 'money', 'money', 'num', 'money', 'money']
    pv = '=IF(E{r}>0,E{r}*(1-(1+{ix})^(-{Y}))/((1+{ix})^(1/12)-1)-(B{r}+D{r}),"")'
    for n, a, h, l, r_ in [('堆高機操作技術士', 9000, 18, 3000, 3000), ('電銲技術士（乙級）', 30000, 120, 12000, 8000), ('移動式起重機操作', 25000, 38, 6000, 6000)]:
        s.row([n, a, h, l, r_, '=IF(E{r}>0,(B{r}+D{r})/E{r},"")', '=IF(E{r}>0,E{r}*12*{Y}-(B{r}+D{r}),"")', pv], k)
    for _ in range(4):
        s.row([None] * 5 + ['=IF(E{r}>0,(B{r}+D{r})/E{r},"")', '=IF(E{r}>0,E{r}*12*{Y}-(B{r}+D{r}),"")', pv], k)
    return wb


# ======================================================================
def business():
    wb = Workbook()
    intro(wb, 'Ken 企業公司試算表', '核心需求：現金流健康、利潤極大化、營運成本控管與碳費衝擊。', [
        ('滾動現金流', '未來 6 個月現金流、Runway 與壓力測試（營收下滑、收款延遲、突發支出）。'),
        ('損益兩平', '固定與變動成本推算最低銷量，含售價敏感度。'),
        ('多維度損益', '各產品線／部門的毛利與營業利益，找出隱形虧損。'),
        ('員工產值', '人均產值、人事成本率與 KPI 掛鉤的獎金分配。'),
        ('碳費試算', '碳費、減量效果與節能投資回收。')])

    s = Sheet(wb, '滾動現金流', '滾動式現金流量預測與壓力測試', '每月月初更新：已發生月份換成實際數字，再往後補一個月。', (24, 16, 16, 16, 16, 16, 16))
    s.inp('open', '期初現金', 3000000)
    s.inp('safe', '安全水位（最低現金）', 1000000)
    s.inp('drop', '壓力測試：營收下降', 0, 'pct')
    s.inp('delay', '壓力測試：收款延遲月數（0–3）', 0, 'int')
    s.inp('shock', '壓力測試：第 1 個月突發支出', 0)
    s.inflation()
    s.inp('cu', '成本隨通膨上升（1＝是，0＝否）', 1, 'int')
    s.r += 1
    months = ['11 月', '12 月', '1 月', '2 月', '3 月', '4 月']
    hr = s.r
    s.header(['項目'] + months)
    labels = ['應收帳款入帳（原定）', '其他收入', '應付帳款', '薪資', '租金與固定開銷', '貸款還款']
    data = [[1800000, 2000000, 1600000, 1200000, 1900000, 2000000], [50000] * 6, [900000, 1000000, 800000, 700000, 950000, 1000000], [650000] * 6, [220000] * 6, [80000] * 6]
    base = s.r
    for l, d in zip(labels, data):
        s.row([l] + d, ['text'] + ['money'] * 6)
    R = {k: base + i for i, k in enumerate(['ar', 'oi', 'ap', 'pay', 'fix', 'loan'])}
    cols = [get_column_letter(c) for c in range(2, 8)]
    calc = [('實際入帳（含延遲與下滑）', lambda c, j: f'=(IF({j}-{{delay}}>=1,INDEX($B${R["ar"]}:$G${R["ar"]},{j}-{{delay}}),0)+{c}{R["oi"]})*(1-{{drop}})'),
            ('流出合計', lambda c, j: f'=({c}{R["ap"]}+{c}{R["pay"]}+{c}{R["fix"]})*IF({{cu}}=1,(1+{{ix}})^(({j}-1)/12),1)+{c}{R["loan"]}' + ('+{shock}' if j == 1 else '')),
            ('淨現金流', None), ('期末現金', None), ('狀態', None)]
    rin = s.r
    s.cell(s.r, 1, calc[0][0], bold=True)
    for j, c in enumerate(cols, 1):
        s.cell(s.r, j + 1, s.ref(calc[0][1](c, j)), 'money', OUT_FILL)
    s.r += 1
    rout = s.r
    s.cell(s.r, 1, calc[1][0], bold=True)
    for j, c in enumerate(cols, 1):
        s.cell(s.r, j + 1, s.ref(calc[1][1](c, j)), 'money', OUT_FILL)
    s.r += 1
    rnet = s.r
    s.cell(s.r, 1, '淨現金流', bold=True)
    for c in cols:
        s.cell(s.r, cols.index(c) + 2, f'={c}{rin}-{c}{rout}', 'money', OUT_FILL)
    s.r += 1
    rend = s.r
    s.cell(s.r, 1, '期末現金', bold=True)
    for j, c in enumerate(cols):
        prev = s.a['open'] if j == 0 else f'{cols[j - 1]}{rend}'
        s.cell(s.r, j + 2, f'={prev}+{c}{rnet}', 'money', OUT_FILL, True)
    s.r += 1
    s.cell(s.r, 1, '狀態', bold=True)
    for c in cols:
        s.cell(s.r, cols.index(c) + 2, f'=IF({c}{rend}<0,"⚠ 資金斷鏈",IF({c}{rend}<{s.a["safe"]},"⚠ 低於水位","正常"))', 'text', OUT_FILL)
    s.ws.conditional_formatting.add(f'B{rend}:G{rend}', FormulaRule(formula=[f'B{rend}<{s.a["safe"]}'], fill=PatternFill('solid', fgColor='F8D3C9')))
    s.r += 2
    s.out('end', '6 個月後期末現金', f'=G{rend}')
    s.out('runway', 'Runway（以虧損月平均估，月）', f'=IFERROR(IF(SUMIF(B{rnet}:G{rnet},"<0")<0,MAX(0,G{rend})/(-SUMIF(B{rnet}:G{rnet},"<0")/6),"現金流為正"),"現金流為正")', 'num')

    s = Sheet(wb, '損益兩平', '專案損益兩平（Break-Even）', '損益兩平銷量＝固定成本 ÷（單價 − 單位變動成本）。', (34, 18, 18, 20, 20))
    s.inp('F', '固定成本（開發、場地、人事、廣告）', 500000)
    s.inp('P', '單價', 1200)
    s.inp('V', '單位變動成本', 700)
    s.inp('T', '目標利潤', 200000)
    s.inp('Q', '預估銷量', 1200, 'int')
    s.inflation()
    s.out('cm', '單位邊際貢獻', '={P}-{V}')
    s.out('cmr', '邊際貢獻率', '=IF({P}>0,{cm}/{P},0)', 'pct')
    s.out('be', '損益兩平銷量', '=IF({cm}>0,ROUNDUP({F}/{cm},0),"售價低於變動成本")', 'int')
    s.out('ber', '損益兩平營收', '=IF({cm}>0,{be}*{P},0)')
    s.out('tq', '達成目標利潤需賣', '=IF({cm}>0,ROUNDUP(({F}+{T})/{cm},0),"—")', 'int')
    s.out('pq', '預估銷量下的利潤', '={Q}*{cm}-{F}')
    s.out('be2', '明年變動成本隨通膨上漲後的損益兩平銷量', '=IF({P}-{V}*(1+{ix})>0,ROUNDUP({F}/({P}-{V}*(1+{ix})),0),"虧損")', 'int')
    s.section('售價敏感度')
    s.header(['售價變動', '售價', '邊際貢獻', '損益兩平銷量', '預估銷量利潤'])
    for d in (-0.2, -0.1, 0, 0.1, 0.2):
        s.row([d, '={P}*(1+A{r})', '=B{r}-{V}', '=IF(C{r}>0,ROUNDUP({F}/C{r},0),"—")', '={Q}*C{r}-{F}'], ['pct', 'money', 'money', 'int', 'money'])

    s = Sheet(wb, '多維度損益', '多維度損益表（P&L）', '毛利為正、營業利益為負的產品線，通常是費用或分攤過重。', (24, 16, 16, 16, 16, 14, 16, 14, 12, 14))
    st = s.header(['產品線／部門', '營收', '營業成本', '直接營業費用', '分攤管銷', '毛利', '毛利率', '營業利益', '營業利益率', '營收占比'])
    rows = [('A 產品線（主力）', 12000000, 6500000, 1800000, 1200000), ('B 產品線（新品）', 3000000, 2100000, 900000, 300000), ('C 服務（維修保養）', 2500000, 800000, 500000, 250000), ('D 電商通路', 4000000, 2600000, 1300000, 400000)] + [(None,) * 5] * 4
    en = st + len(rows) - 1
    k = ['text', 'money', 'money', 'money', 'money', 'money', 'pct', 'money', 'pct', 'pct']
    for row in rows:
        s.row(list(row) + ['=IF(B{r}="","",B{r}-C{r})', '=IF(N(B{r})>0,F{r}/B{r},"")', '=IF(B{r}="","",F{r}-D{r}-E{r})', '=IF(N(B{r})>0,H{r}/B{r},"")', f'=IF(N(B{{r}})>0,B{{r}}/SUM(B{st}:B{en}),"")'], k)
    s.row(['合計', f'=SUM(B{st}:B{en})', f'=SUM(C{st}:C{en})', f'=SUM(D{st}:D{en})', f'=SUM(E{st}:E{en})', f'=SUM(F{st}:F{en})', '=IF(B{r}>0,F{r}/B{r},0)', f'=SUM(H{st}:H{en})', '=IF(B{r}>0,H{r}/B{r},0)', 1], k)
    s.ws.conditional_formatting.add(f'H{st}:H{en}', FormulaRule(formula=[f'AND(ISNUMBER(H{st}),H{st}<0)'], fill=PatternFill('solid', fgColor='F8D3C9')))
    s.note('紅底＝營業利益為負的隱形虧損。')

    s = Sheet(wb, '員工產值', '員工產值與獎金分配', '分配權重＝（貢獻 − 人事成本）× KPI 達成率。', (18, 10, 18, 18, 12, 16, 14, 16, 16))
    s.inp('profit', '公司稅前淨利', 10000000)
    s.inp('pool', '獎金提撥比例', 0.15, 'pct')
    s.inflation()
    s.out('P', '獎金池', '={profit}*{pool}')
    s.out('raise', '明年維持實質薪資需調薪', '={ix}', 'pct')
    s.r += 1
    st = s.header(['部門', '人數', '營收或毛利貢獻', '年人事成本', 'KPI 達成率', '人均產值', '人事成本率', '部門獎金', '人均獎金'])
    rows = [('業務部', 8, 30000000, 8000000, 1.1), ('產品研發', 6, 12000000, 7200000, .95), ('客服營運', 5, 5000000, 3500000, 1), ('行政管理', 3, 2000000, 2400000, .9)] + [(None,) * 5] * 3
    en = st + len(rows) - 1
    w = f'SUMPRODUCT((C{st}:C{en}-D{st}:D{en}>0)*(C{st}:C{en}-D{st}:D{en})*E{st}:E{en})'
    for row in rows:
        s.row(list(row) + ['=IF(N(B{r})>0,C{r}/B{r},"")', '=IF(N(C{r})>0,D{r}/C{r},"")', f'=IF(B{{r}}="","",IF({w}>0,{{P}}*MAX(0,C{{r}}-D{{r}})*E{{r}}/{w},0))', '=IF(N(B{r})>0,H{r}/B{r},"")'],
              ['text', 'int', 'money', 'money', 'pct', 'money', 'pct', 'money', 'money'])

    s = Sheet(wb, '碳費試算', '碳費與綠色轉型成本', '依環境部碳費制度：2025 年排放量起算，2026 年 5 月首次繳納；費率與條件以最新公告為準。', (44, 20, 44))
    s.inp('E', '年排放量（公噸 CO2e，直接＋間接）', 40000, 'int')
    s.inp('type', '適用費率（一般／優惠A／優惠B）', '一般', 'text')
    s.validate_list(s.a['type'].replace('$', ''), ['一般', '優惠A', '優惠B'])
    s.inp('leak', '高碳洩漏風險係數（一般 1，適用者 0.2）', 1, 'num')
    s.inp('cut', '減量計畫：減少排放', 0.10, 'pct')
    s.inp('cap', '節能設備投資', 20000000)
    s.inp('sv', '每年節省能源費用', 3000000)
    s.inflation()
    s.out('price', '費率（元／公噸）', '=IF({type}="優惠A",50,IF({type}="優惠B",100,300))', 'int')
    fee = 'IF({X}<25000,0,IF({price}=300,MAX(0,{X}-25000),{X})*{leak}*{price})'
    s.out('now', '目前年碳費', '=' + fee.replace('{X}', '{E}'), 'money', '一般費率可扣除 2.5 萬公噸；優惠費率不得扣除')
    s.out('after', '減量後年碳費', '=' + fee.replace('{X}', '({E}*(1-{cut}))'))
    s.out('year1', '第一年總節省（能源費隨通膨＋碳費）', '={sv}*(1+{ix})+{now}-{after}')
    s.out('pb', '設備投資回收年數', '=IF({year1}>0,{cap}/{year1},"無法回收")', 'num')
    return wb


# ======================================================================
def sales():
    wb = Workbook()
    intro(wb, 'Ken 直銷經營試算表', '核心需求：獎金試算、客戶回購、庫存與行銷成本。本表為通用模板，獎金請依所屬公司正式制度填寫；不保證任何收入。', [
        ('獎金晉升', '依聘階門檻推算本月獎金與晉升差額。'),
        ('回購CRM', '產品快用完自動亮紅燈，並計算顧客終身價值（LTV）。'),
        ('進銷存', '庫存、進貨成本、利潤，獨立記錄試用品成本。'),
        ('行銷ROI', '各管道的名單成本、轉換率、獲客成本與 LTV／CAC。'),
        ('組織成長', '加入留存率與活躍率的組織成長現實檢查。')])

    s = Sheet(wb, '獎金晉升', '階梯式獎金與晉升差額', '聘階由低到高排列；以「團隊業績 × 聘階比例」簡化估算，雙軌制、太陽線、代數獎金請依公司辦法調整。', (22, 18, 18, 14, 18))
    s.inp('p', '本月個人業績', 30000)
    s.inflation()
    s.section('下線團隊本月業績')
    t0 = s.r
    for n, v in [('夥伴 A', 60000), ('夥伴 B', 45000), ('夥伴 C', 30000), ('夥伴 D', 25000)] + [(None, None)] * 8:
        s.row([n, v], ['text', 'money'])
    t1 = s.r - 1
    s.out('T', '團隊總業績（含個人）', f'={{p}}+SUM(B{t0}:B{t1})')
    s.section('聘階制度（由低到高）')
    r0 = s.header(['聘階', '個人業績門檻', '團隊業績門檻', '獎金比例', '是否達成'])
    for n, a, b, c in [('會員', 0, 0, .03), ('銀級', 20000, 100000, .06), ('金級', 30000, 300000, .09), ('鑽石', 40000, 800000, .12), (None, None, None, None)]:
        s.row([n, a, b, c, '=IF(A{r}="","",IF(AND({p}>=B{r},{T}>=C{r}),"✔","—"))'], ['text', 'money', 'money', 'pct', 'text'])
    r1 = s.r - 1
    s.r += 1
    s.out('rank', '目前聘階', f'=IFERROR(LOOKUP(2,1/(E{r0}:E{r1}="✔"),A{r0}:A{r1}),"未達任何聘階")', 'text')
    s.out('bonus', '本月預估獎金', f'=IFERROR({{T}}*LOOKUP(2,1/(E{r0}:E{r1}="✔"),D{r0}:D{r1}),0)')
    s.out('nextn', '下一聘階', f'=IFERROR(INDEX(A{r0}:A{r1},MATCH("—",E{r0}:E{r1},0)),"已達最高")', 'text')
    s.out('gp', '晉升還差：個人業績', f'=IFERROR(MAX(0,INDEX(B{r0}:B{r1},MATCH("—",E{r0}:E{r1},0))-{{p}}),0)')
    s.out('gt', '晉升還差：團隊業績', f'=IFERROR(MAX(0,INDEX(C{r0}:C{r1},MATCH("—",E{r0}:E{r1},0))-{{T}}),0)')
    s.out('real', '獎金一年後實質購買力', '={bonus}/(1+{ix})')

    s = Sheet(wb, '回購CRM', '客戶回購週期與 LTV', '預計用完日＝購買日＋數量 × 每份可用天數；在提醒日當天自動亮紅燈。客戶個資請妥善保存。', (14, 22, 14, 10, 14, 14, 14, 16))
    s.inp('rem', '提前幾天提醒', 5, 'int')
    st = s.header(['客戶', '產品', '購買日', '數量', '每份可用天數', '預計用完', '剩餘天數', '燈號'])
    k = ['text', 'text', 'date', 'int', 'int', 'date', 'int', 'text']
    tail = ['=IF(C{r}="","",C{r}+D{r}*E{r})', '=IF(C{r}="","",F{r}-TODAY())', '=IF(C{r}="","",IF(G{r}<={rem},"🔴 該聯絡",IF(G{r}<={rem}+7,"🟡 一週內","🟢 正常")))']
    for n, p, d_, q, dd in [('王姐', '保健食品 A', date(2026, 9, 20), 1, 30), ('陳大哥', '保養組', date(2026, 8, 30), 2, 30), ('林阿姨', '保健食品 B', date(2026, 10, 1), 1, 20), ('張小姐', '洗沐組', date(2026, 7, 15), 3, 30)]:
        s.row([n, p, d_, q, dd] + tail, k)
    for _ in range(46):
        s.row([None] * 5 + tail, k)
    en = s.r - 1
    s.ws.conditional_formatting.add(f'A{st}:H{en}', FormulaRule(formula=[f'AND($C{st}<>"",$G{st}<={s.a["rem"]})'], fill=PatternFill('solid', fgColor='F8D3C9')))
    s.ws.freeze_panes = f'A{st}'
    s.r += 1
    s.section('顧客終身價值（LTV）')
    s.inp('avg', '平均客單價', 3600)
    s.inp('t', '每年回購次數', 6, 'int')
    s.inp('y', '預期維繫年數', 3, 'int')
    s.inp('m', '毛利率', 0.30, 'pct')
    s.inp('c', '每位客戶每年維繫成本（送禮、咖啡）', 1500)
    s.out('ltv', '顧客終身毛利（LTV）', '={avg}*{t}*{y}*{m}')
    s.out('cost', '維繫成本', '={c}*{y}')
    s.out('roi', '維繫投資報酬率', '=IF({cost}>0,({ltv}-{cost})/{cost},0)', 'pct')
    s.out('red', '今天該聯絡人數', f'=COUNTIF(H{st}:H{en},"🔴 該聯絡")', 'int')

    s = Sheet(wb, '進銷存', '進銷存與真實利潤（含試用品）', '試用品是行銷費用，獨立記錄才能看出真實利潤。期末庫存為負代表紀錄有誤。', (20, 10, 10, 12, 10, 12, 10, 10, 14, 14, 14, 14, 14))
    st = s.header(['品項', '期初庫存', '本期進貨', '進貨單價', '售出', '售價', '送出試用', '期末庫存', '銷售額', '銷貨成本', '試用品成本', '真實利潤', '庫存金額'])
    k = ['text', 'int', 'int', 'money', 'int', 'money', 'int', 'int', 'money', 'money', 'money', 'money', 'money']
    tail = ['=IF(A{r}="","",B{r}+C{r}-E{r}-G{r})', '=IF(A{r}="","",E{r}*F{r})', '=IF(A{r}="","",E{r}*D{r})', '=IF(A{r}="","",G{r}*D{r})', '=IF(A{r}="","",I{r}-J{r}-K{r})', '=IF(A{r}="","",H{r}*D{r})']
    for row in [('保健食品 A', 5, 10, 1200, 9, 1680, 2), ('保養組', 3, 6, 2100, 5, 2980, 1), ('體驗包', 0, 20, 150, 0, 0, 18)]:
        s.row(list(row) + tail, k)
    for _ in range(12):
        s.row([None] * 7 + tail, k)
    en = s.r - 1
    s.r += 1
    s.out('rev', '銷售額', f'=SUM(I{st}:I{en})')
    s.out('smp', '試用品成本', f'=SUM(K{st}:K{en})')
    s.out('pro', '真實利潤（扣試用品）', f'=SUM(L{st}:L{en})')
    s.out('stk', '庫存積壓金額', f'=SUM(M{st}:M{en})')

    s = Sheet(wb, '行銷ROI', '個人品牌行銷 ROI', 'LTV ÷ CAC 小於 1 代表越做越賠；3 倍以上通常才算健康。', (20, 14, 10, 10, 14, 16, 14, 12, 14, 14, 12))
    st = s.header(['管道', '花費', '名單數', '成交數', '平均首單', '預期終身毛利', '名單成本', '轉換率', '獲客成本', '首單 ROI', 'LTV÷CAC'])
    k = ['text', 'money', 'int', 'int', 'money', 'money', 'money', 'pct', 'money', 'pct', 'num']
    tail = ['=IF(N(C{r})>0,B{r}/C{r},"")', '=IF(N(C{r})>0,D{r}/C{r},"")', '=IF(N(D{r})>0,B{r}/D{r},"")', '=IF(N(B{r})>0,(D{r}*E{r}-B{r})/B{r},"")', '=IF(N(B{r})>0,D{r}*F{r}/B{r},"")']
    for row in [('FB／IG 廣告', 15000, 60, 4, 3600, 6480), ('線下說明會', 12000, 25, 5, 4200, 7500), ('試用包', 6000, 30, 3, 2500, 4500), ('朋友轉介紹', 2000, 8, 3, 3600, 6480)]:
        s.row(list(row) + tail, k)
    for _ in range(4):
        s.row([None] * 6 + tail, k)
    en = s.r - 1
    s.r += 1
    s.out('sp', '總花費', f'=SUM(B{st}:B{en})')
    s.out('cac', '平均獲客成本（CAC）', f'=IF(SUM(D{st}:D{en})>0,{{sp}}/SUM(D{st}:D{en}),0)')

    s = Sheet(wb, '組織成長', '組織成長現實檢查', '「每人找 2 人、每月倍增」只是理想；加入留存率與活躍率後的估算。直銷收入應來自商品實際銷售，不保證收入。', (12, 16, 16, 18, 16, 18))
    s.inp('n0', '起始團隊人數（含自己）', 1, 'int')
    s.inp('inv', '每位活躍夥伴每月帶進', 1, 'num')
    s.inp('keep', '每月留存率', 0.70, 'pct')
    s.inp('act', '活躍率', 0.40, 'pct')
    s.inp('sal', '每位活躍者月業績', 3000)
    s.inp('rate', '團隊獎金比例', 0.08, 'pct')
    s.inp('hrs', '每月投入時數', 40, 'int')
    s.r += 1
    st = s.header(['月份', '團隊人數', '活躍人數', '團隊業績', '預估獎金', '理想倍增人數'])
    for m in range(1, 25):
        p = s.r - 1
        prevN = '{n0}' if m == 1 else f'B{p}'
        prevI = '{n0}' if m == 1 else f'F{p}'
        s.row([m, f'={prevN}*{{keep}}+MAX(1,{prevN}*{{act}})*{{inv}}', '=MAX(1,B{r}*{act})', '=C{r}*{sal}', '=D{r}*{rate}', f'={prevI}*2'], ['int', 'num', 'num', 'money', 'money', 'int'])
    en = s.r - 1
    s.r += 1
    s.out('b12', '第 12 個月預估獎金', f'=E{st + 11}')
    s.out('hr', '換算時薪', f'=IF({{hrs}}>0,E{st + 11}/{{hrs}},0)', 'money', '對照 2026 年基本工資時薪 196 元')
    s.out('pt', '同樣時數做時薪兼差', '={hrs}*196')
    return wb


def main():
    try:
        OUT.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        sys.exit(f'建立 {OUT} 失敗：{e}')
    for name, build in [('ken-basics.xlsx', basics), ('ken-seniors.xlsx', seniors), ('ken-white-collar.xlsx', white),
                        ('ken-blue-collar.xlsx', blue), ('ken-business.xlsx', business), ('ken-direct-sales.xlsx', sales)]:
        wb = build()
        for ws in wb.worksheets:
            ws.sheet_view.zoomScale = 110
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.fitToWidth = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToHeight = 0
        wb.properties.creator = 'AK Wealth Global'
        wb.properties.title = name
        try:
            wb.save(OUT / name)
        except OSError as e:
            sys.exit(f'寫入 {OUT / name} 失敗：{e}')
        print('wrote', OUT / name)


if __name__ == '__main__':
    main()
